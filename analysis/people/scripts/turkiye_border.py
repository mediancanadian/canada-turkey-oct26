import openpyxl,xlrd,json
O='../people/'
wb=openpyxl.load_workbook(O+'raw/turkiye/sinir_2025.xlsx',read_only=True); ws=wb['C-Mil Göre G.Yabancı']
rows=[list(r) for r in ws.iter_rows(values_only=True)]
hdr=[r for r in rows[:5] if r and any(c=='MİLLİYET' for c in r)][0]
off=hdr.index('MİLLİYET'); years=[c for c in hdr[off+1:] if isinstance(c,int)]
data={}
for r in rows[3:]:
    if not r or len(r)<=off: continue
    name=r[off] if r[off] is not None else (r[off+1] if len(r)>off+1 and isinstance(r[off+1],str) else None)
    if not isinstance(name,str): continue
    vals=[c for c in r if isinstance(c,(int,float))]
    if len(vals)==len(years): data[name.strip()]=dict(zip(years,vals))
can=data['Kanada']; print('Kanada',can)
tot=[k for k in data if 'TOPLAM' in k.upper() or 'GENEL' in k.upper()]; print('totals',tot,[data[k][2025] for k in tot])
skip=set(tot)|{k for k in data if k.upper()==k and len(k)>4}
rk={}
for y in years:
    s=sorted([(v[y],k) for k,v in data.items() if k not in skip],reverse=True)
    rk[y]=[k for _,k in s].index('Kanada')+1
print('rank',rk)
# T7 monthly 2025
ws=wb['T7-Mil-Ay Göre G.Yabancı']
m25=[[c for c in r if c is not None] for r in ws.iter_rows(values_only=True) if r and 'Kanada' in [str(c).strip() for c in r if c is not None]][0]
print('2025 monthly',m25)
# 2026 Aug bulletin
bk=xlrd.open_workbook(O+'raw/turkiye/sinir_2026aug.xls')
hits={}
for sh in bk.sheets():
    for i in range(sh.nrows):
        row=sh.row_values(i)
        if any(str(c).strip()=='Kanada' for c in row):
            hdrs=[sh.row_values(j) for j in range(max(0,i-200),min(8,sh.nrows))]
            hits[sh.name]=row; print('2026',sh.name,i,[c for c in row if c!=''][:20]); break
heads={}
for sh in bk.sheets():
    if sh.name in hits:
        heads[sh.name]=[[c for c in sh.row_values(j) if c!=''] for j in range(0,6)]
        print(sh.name,'HEAD',heads[sh.name])
json.dump({'annual_arrivals_canadian_nationals':can,'rank_among_nationalities':rk,'monthly_2025':m25,'bulletin_2026_aug_rows':hits,'bulletin_2026_aug_heads':heads},open(O+'outputs/turkiye_border_canada.json','w'),indent=1,ensure_ascii=False,default=str)
